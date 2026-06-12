import argparse
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "docs" / "作业验收预期表_策略B_时间顺序_修正版.xlsx"
FALLBACK_XLSX = ROOT / "docs" / "作业验收预期表_策略B_时间顺序.xlsx"
OUT_JSON = ROOT / "docs" / "strategy_b_verification.json"

EVENTS = [
    ["06:00", "A", "V1", "T", "40"], ["06:05", "A", "V2", "T", "30"],
    ["06:10", "A", "V3", "F", "60"], ["06:20", "A", "V2", "O", "0"],
    ["06:25", "A", "V4", "T", "20"], ["06:30", "A", "V5", "T", "20"],
    ["06:40", "A", "V6", "T", "20"], ["06:50", "A", "V7", "T", "10"],
    ["07:00", "A", "V8", "F", "90"], ["07:10", "A", "V9", "F", "30"],
    ["07:15", "A", "V10", "T", "10"], ["07:20", "A", "V11", "F", "60"],
    ["07:25", "A", "V12", "T", "10"], ["07:30", "A", "V13", "T", "7.5"],
    ["07:35", "A", "V14", "F", "75"], ["07:40", "A", "V15", "F", "45"],
    ["08:00", "A", "V16", "T", "5"], ["08:20", "A", "V17", "T", "15"],
    ["08:30", "A", "V18", "T", "20"], ["08:35", "A", "V19", "T", "25"],
    ["09:00", "A", "V20", "F", "30"], ["09:10", "A", "V7", "O", "0"],
    ["09:20", "A", "V11", "O", "0"], ["09:30", "A", "V18", "O", "0"],
    ["09:35", "A", "V20", "O", "0"], ["09:50", "A", "V21", "F", "30"],
    ["10:00", "A", "V22", "T", "10"], ["10:05", "C", "V19", "F", "25"],
    ["10:10", "C", "V21", "F", "10"], ["10:20", "C", "V22", "F", "10"],
    ["10:30", "B", "T1", "O", "60"], ["10:50", "B", "F1", "O", "120"],
]

PILE_IDS = ["F1", "F2", "F3", "T1", "T2"]
TOKEN = None


def minutes(hhmm):
    h, m = [int(x) for x in hhmm.split(":")]
    return h * 60 + m


def http(base_url, method, path, body=None, token=True, timeout=5):
    data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    headers = {"Accept": "application/json"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    if token and TOKEN:
        headers["Authorization"] = "Bearer " + TOKEN
    request = urllib.request.Request(base_url + path, data=data, headers=headers, method=method)
    with urllib.request.urlopen(request, timeout=timeout) as response:
        payload = json.loads(response.read().decode("utf-8"))
    if not payload.get("success", False):
        raise RuntimeError(payload.get("message") or f"{method} {path} failed")
    return payload.get("data")


def try_http(base_url, method, path, body=None, token=True):
    try:
        return http(base_url, method, path, body, token), None
    except Exception as exc:
        return None, exc


def ensure_auth(base_url):
    global TOKEN
    http(base_url, "GET", "/api/health", token=False)
    _, err = try_http(base_url, "GET", "/api/admin/snapshot", token=False)
    if err is None:
        return
    if isinstance(err, urllib.error.HTTPError) and err.code in (401, 403):
        username = "verify_strategy_b_admin"
        password = "verify_strategy_b_123456"
        try_http(base_url, "POST", "/api/auth/register", {
            "username": username, "password": password, "role": "ADMIN"
        }, token=False)
        auth = http(base_url, "POST", "/api/auth/login", {
            "username": username, "password": password
        }, token=False)
        TOKEN = auth.get("token")
        http(base_url, "GET", "/api/admin/snapshot")
        return
    raise RuntimeError(f"cannot access {base_url}: {err}")


def reset_station(base_url):
    http(base_url, "POST", "/api/admin/config", {
        "waitingCapacity": 10,
        "pileSlotCapacity": 3,
        "fastPileCount": 3,
        "slowPileCount": 2,
        "parkingRate": 0.5,
        "parkingGracePeriodMinutes": 5,
    })
    http(base_url, "POST", "/api/admin/clock/reset?time=06:00")


def run_events(base_url, include_strategy=True):
    reset_station(base_url)
    snapshots = []
    for t, action, subject, mode, amount in EVENTS:
        body = {"time": t, "action": action, "subject": subject, "mode": mode, "amount": amount}
        if include_strategy:
            body["strategy"] = "TIME_ORDER"
        snapshots.append(http(base_url, "POST", "/api/acceptance/event", body, token=False))
    return snapshots


def settle_for_details(base_url, minutes_to_advance=300):
    http(base_url, "POST", f"/api/admin/clock/advance?minutes={minutes_to_advance}")
    return http(base_url, "GET", "/api/admin/snapshot")


def dec(value):
    return Decimal(str(value or 0))


def fmt(value, places=None):
    d = dec(value)
    if places is not None:
        d = d.quantize(Decimal("1." + "0" * places), rounding=ROUND_HALF_UP)
    text = format(d.normalize(), "f")
    return "0" if text == "-0" else text


def hhmm(value):
    if not value:
        return None
    match = re.search(r"(\d{2}:\d{2})", str(value))
    return match.group(1) if match else None


def rate_at(hhmm_text):
    t = minutes(hhmm_text)
    if 600 <= t < 900 or 1080 <= t < 1260:
        return Decimal("1.0")
    if 420 <= t < 600 or 900 <= t < 1080 or 1260 <= t < 1380:
        return Decimal("0.7")
    return Decimal("0.4")


def current_fee(req, snap):
    start = hhmm(req.get("startedAt"))
    end = hhmm(snap.get("now"))
    kwh = dec(req.get("chargedKwh"))
    if not start or not end or kwh <= 0:
        return Decimal("0")
    total = max(1, minutes(end) - minutes(start))
    cursor = minutes(start)
    charge_fee = Decimal("0")
    while cursor < minutes(start) + total:
        next_cursor = minutes(start) + total
        for boundary in [420, 600, 900, 1080, 1260, 1380]:
            if cursor < boundary <= minutes(start) + total:
                next_cursor = min(next_cursor, boundary)
                break
        part = Decimal(next_cursor - cursor) / Decimal(total)
        label = f"{(cursor // 60) % 24:02d}:{cursor % 60:02d}"
        charge_fee += kwh * part * rate_at(label)
        cursor = next_cursor
    return charge_fee + kwh * Decimal("0.8")


def latest_request_by_vehicle(snap):
    result = {}
    for req in snap.get("requests", []):
        old = result.get(req.get("vehicleId"))
        if not old or str(req.get("createdAt")) >= str(old.get("createdAt")):
            result[req.get("vehicleId")] = req
    return result


def pile_cell(snap, pile_id, slot):
    pile = next((p for p in snap.get("piles", []) if p.get("id") == pile_id), None)
    if not pile:
        return ""
    if pile.get("status") == "FAULT":
        return "✕故障"
    if slot == 0:
        vid = pile.get("currentVehicle")
        if not vid:
            return ""
        req = latest_request_by_vehicle(snap).get(vid, {})
        return f"({vid},{fmt(req.get('chargedKwh'), 2)},{fmt(current_fee(req, snap), 2)})"
    queued = pile.get("queuedVehicles") or []
    if len(queued) < slot:
        return ""
    return f"({queued[slot - 1]},0,0)"


def waiting_cell(snap):
    reqs = latest_request_by_vehicle(snap)
    values = []
    for vid in (snap.get("waitingFast") or []) + (snap.get("waitingSlow") or []):
        req = reqs.get(vid, {})
        mode = "F" if req.get("mode") == "FAST" else "T"
        values.append(f"({vid},{mode},{fmt(req.get('requestedKwh'))})")
    return " ".join(values)


def redispatch_cell(snap):
    values = []
    for req in snap.get("requests", []):
        if req.get("status") == "REDISPATCHING":
            mode = "F" if req.get("mode") == "FAST" else "T"
            values.append(f"({req.get('vehicleId')},{mode},{fmt(req.get('requestedKwh'))})")
    return " ".join(values)


def actual_table_rows(snapshots):
    rows = [
        ["", "", "(车号,已充电量,当前费用)", "", "", "", "", "(车号,充电类型,充电量)", "故障后车辆，不属于普通等候区"],
        ["时刻", "事件 (类型,id,模式,值)", "快充1 (F1)", "快充2 (F2)", "快充3 (F3)", "慢充1 (T1)", "慢充2 (T2)", "等候区(10辆)", "故障重调度队列"],
    ]
    for event, snap in zip(EVENTS, snapshots):
        rows.append([
            event[0], f"({event[1]},{event[2]},{event[3]},{event[4]})",
            *[pile_cell(snap, pile_id, 0) for pile_id in PILE_IDS],
            waiting_cell(snap), redispatch_cell(snap),
        ])
        rows.append(["", "", *[pile_cell(snap, pile_id, 1) for pile_id in PILE_IDS], "", ""])
        rows.append(["", "", *[pile_cell(snap, pile_id, 2) for pile_id in PILE_IDS], "", ""])
    return rows


def excel_table_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["预期表"]
    return [[ws.cell(r, c).value for c in range(1, 10)] for r in range(1, 1 + 2 + len(EVENTS) * 3)]


def excel_detail_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["详单(自动结算)"]
    rows = []
    for r in range(1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, 10)]
        if any(v is not None for v in values):
            rows.append(values)
    return rows


def actual_detail_rows(snap):
    by_request = {r.get("id"): r for r in snap.get("requests", [])}
    rows = []
    for detail in snap.get("details", []):
        req = by_request.get(detail.get("requestId"), {})
        rows.append([
            detail.get("vehicleId"), detail.get("pileId"),
            hhmm(detail.get("startedAt")), hhmm(detail.get("stoppedAt")),
            fmt(detail.get("chargedKwh"), 2), fmt(detail.get("chargeFee"), 2),
            fmt(detail.get("serviceFee"), 2), fmt(detail.get("totalFee"), 2),
            detail.get("status") or req.get("status", ""),
        ])
    return rows


def normalize_cell(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).replace("X故障", "✕故障")


def compare_rows(expected, actual):
    mismatches = []
    for r, (erow, arow) in enumerate(zip(expected, actual), start=1):
        for c, (e, a) in enumerate(zip(erow, arow), start=1):
            ne, na = normalize_cell(e), normalize_cell(a)
            if ne != na:
                mismatches.append({"row": r, "col": c, "expected": ne, "actual": na})
    return mismatches


def normalize_detail(rows):
    clean = []
    for row in rows:
        if row and normalize_cell(row[0]) in {"车辆", "车辆编号", "vehicleId", "车号"}:
            continue
        clean.append([normalize_cell(v) for v in row])
    return clean


def compare_details(expected, actual):
    exp = normalize_detail(expected)
    act = normalize_detail(actual)
    mismatches = []
    for i in range(max(len(exp), len(act))):
        erow = exp[i] if i < len(exp) else []
        arow = act[i] if i < len(act) else []
        if erow != arow:
            mismatches.append({"row": i + 1, "expected": erow, "actual": arow})
    return mismatches


def vehicles_with_status(snap, status, mode=None):
    values = []
    for req in snap.get("requests", []):
        if req.get("status") == status and (mode is None or req.get("mode") == mode):
            values.append(req.get("vehicleId"))
    return values


def pile_by_id(snap, pile_id):
    return next(p for p in snap.get("piles", []) if p.get("id") == pile_id)


def details_by_vehicle(snap, vehicle_id):
    return [d for d in snap.get("details", []) if d.get("vehicleId") == vehicle_id]


def invariant_checks(snapshots, final_snapshot, default_snapshots=None):
    checks = []

    def check(name, ok, actual=None, expected=None):
        checks.append({"name": name, "ok": bool(ok), "actual": actual, "expected": expected})

    for index, snap in enumerate(snapshots):
        m = snap.get("stationConfig", {}).get("pileSlotCapacity", 3)
        offenders = []
        for pile in snap.get("piles", []):
            occupied = (1 if pile.get("currentVehicle") else 0) + len(pile.get("queuedVehicles") or [])
            if pile.get("status") != "FAULT" and occupied > m:
                offenders.append(f"{pile.get('id')}={occupied}")
        check(f"每桩不超过 M=3 after {EVENTS[index][0]}", not offenders, offenders, "no pile over capacity")

    snap_1030 = snapshots[30]
    check("10:30 T1 故障", pile_by_id(snap_1030, "T1").get("status") == "FAULT",
          pile_by_id(snap_1030, "T1").get("status"), "FAULT")
    check("10:30 T2 桩内为 V12 + V6/V10",
          pile_by_id(snap_1030, "T2").get("currentVehicle") == "V12"
          and pile_by_id(snap_1030, "T2").get("queuedVehicles") == ["V6", "V10"],
          [pile_by_id(snap_1030, "T2").get("currentVehicle"), pile_by_id(snap_1030, "T2").get("queuedVehicles")],
          ["V12", ["V6", "V10"]])
    check("10:30 故障重调度队列为 V13/V16/V17",
          vehicles_with_status(snap_1030, "REDISPATCHING", "SLOW") == ["V13", "V16", "V17"],
          vehicles_with_status(snap_1030, "REDISPATCHING", "SLOW"), ["V13", "V16", "V17"])

    snap_1050 = snapshots[31]
    check("10:50 F1 故障", pile_by_id(snap_1050, "F1").get("status") == "FAULT",
          pile_by_id(snap_1050, "F1").get("status"), "FAULT")
    check("10:50 F2/F3 为 V15/V19",
          pile_by_id(snap_1050, "F2").get("currentVehicle") == "V15"
          and pile_by_id(snap_1050, "F3").get("currentVehicle") == "V19",
          [pile_by_id(snap_1050, "F2").get("currentVehicle"), pile_by_id(snap_1050, "F3").get("currentVehicle")],
          ["V15", "V19"])
    check("10:50 慢充故障重调度仍为 V13/V16/V17",
          vehicles_with_status(snap_1050, "REDISPATCHING", "SLOW") == ["V13", "V16", "V17"],
          vehicles_with_status(snap_1050, "REDISPATCHING", "SLOW"), ["V13", "V16", "V17"])

    v21_details = details_by_vehicle(final_snapshot, "V21")
    check("V21 10:10 变更生成 10 度详单",
          len(v21_details) == 1 and fmt(v21_details[0].get("chargedKwh"), 2) == "10",
          [[d.get("pileId"), hhmm(d.get("startedAt")), hhmm(d.get("stoppedAt")), fmt(d.get("chargedKwh"), 2)] for d in v21_details],
          [["F1", "09:50", "10:10", "10"]])

    v13_details = details_by_vehicle(final_snapshot, "V13")
    check("V13 最终到 T1 12:30-13:15",
          len(v13_details) == 1 and v13_details[0].get("pileId") == "T1"
          and hhmm(v13_details[0].get("startedAt")) == "12:30"
          and hhmm(v13_details[0].get("stoppedAt")) == "13:15",
          [[d.get("pileId"), hhmm(d.get("startedAt")), hhmm(d.get("stoppedAt"))] for d in v13_details],
          [["T1", "12:30", "13:15"]])

    unfinished = [r.get("vehicleId") for r in final_snapshot.get("requests", [])
                  if r.get("status") in {"WAITING_AREA", "REDISPATCHING", "PILE_QUEUE", "CHARGING"}]
    check("推进 300 分钟后没有未完成车辆", not unfinished, unfinished, [])

    if default_snapshots:
        default_1030 = default_snapshots[30]
        check("不传 strategy 时也默认策略B",
              vehicles_with_status(default_1030, "REDISPATCHING", "SLOW") == ["V13", "V16", "V17"],
              vehicles_with_status(default_1030, "REDISPATCHING", "SLOW"), ["V13", "V16", "V17"])

    return checks


def print_result(title, ok):
    print(f"[{'OK' if ok else 'FAIL'}] {title}")


def main():
    parser = argparse.ArgumentParser(description="Verify strategy B acceptance Excel against the running backend.")
    parser.add_argument("--base-url", default="http://localhost:8080")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX if DEFAULT_XLSX.exists() else FALLBACK_XLSX))
    parser.add_argument("--skip-default-check", action="store_true")
    args = parser.parse_args()

    base_url = args.base_url.rstrip("/")
    xlsx = Path(args.xlsx)
    if not xlsx.exists():
        raise FileNotFoundError(xlsx)

    started = time.time()
    ensure_auth(base_url)
    snapshots = run_events(base_url, include_strategy=True)
    final_snapshot = settle_for_details(base_url)

    expected_table = excel_table_rows(xlsx)
    actual_table = actual_table_rows(snapshots)
    table_mismatches = compare_rows(expected_table, actual_table)

    expected_details = excel_detail_rows(xlsx)
    actual_details = actual_detail_rows(final_snapshot)
    detail_mismatches = compare_details(expected_details, actual_details)

    default_snapshots = None
    if not args.skip_default_check:
        default_snapshots = run_events(base_url, include_strategy=False)

    checks = invariant_checks(snapshots, final_snapshot, default_snapshots)
    failed_checks = [c for c in checks if not c["ok"]]

    result = {
        "base_url": base_url,
        "xlsx": str(xlsx),
        "elapsed_seconds": round(time.time() - started, 2),
        "table_mismatches": table_mismatches,
        "detail_mismatches": detail_mismatches,
        "checks": checks,
        "actual_details": actual_details,
    }
    OUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Strategy B verification")
    print(f"Backend: {base_url}")
    print(f"Excel:   {xlsx}")
    print(f"Report:  {OUT_JSON}")
    print_result("预期表主表与后端逐事件快照一致", not table_mismatches)
    print_result("详单(自动结算)与后端最终详单一致", not detail_mismatches)
    print_result("关键不变量全部通过", not failed_checks)

    if table_mismatches:
        print("\n预期表差异前 10 条:")
        for item in table_mismatches[:10]:
            print(json.dumps(item, ensure_ascii=False))
    if detail_mismatches:
        print("\n详单差异:")
        for item in detail_mismatches:
            print(json.dumps(item, ensure_ascii=False))
    if failed_checks:
        print("\n关键不变量失败:")
        for item in failed_checks:
            print(json.dumps(item, ensure_ascii=False))

    if table_mismatches or detail_mismatches or failed_checks:
        sys.exit(1)


if __name__ == "__main__":
    main()
