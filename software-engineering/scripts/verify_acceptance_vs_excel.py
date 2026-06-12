import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT_HTML = DOCS / "web_acceptance_excel_compare.html"
OUT_JSON = DOCS / "web_acceptance_excel_compare.json"
BASE_URL = os.environ.get("EV_VERIFY_BASE_URL", "http://localhost:18082").rstrip("/")
TIMEOUT = float(os.environ.get("EV_VERIFY_TIMEOUT", "2"))

EVENTS = [
    ["06:00", "A", "V1", "T", "40"], ["06:05", "A", "V2", "T", "30"],
    ["06:10", "A", "V3", "F", "60"], ["06:20", "A", "V2", "O", "0"],
    ["06:25", "A", "V4", "T", "20"], ["06:30", "A", "V5", "T", "20"],
    ["06:40", "A", "V6", "T", "20"], ["06:50", "A", "V7", "T", "10"],
    ["07:00", "A", "V8", "F", "90"], ["07:10", "A", "V9", "F", "30"],
    ["07:15", "A", "V10", "T", "10"], ["07:20", "A", "V11", "F", "60"],
    ["07:25", "A", "V12", "T", "10"], ["07:30", "A", "V13", "T", "7.5"],
    ["07:35", "A", "V14", "F", "75"], ["07:40", "A", "V15", "F", "45"],
    ["08:00", "A", "V16", "T", "5"], ["08:20", "A", "V17", "T", "15"],
    ["08:30", "A", "V18", "T", "20"], ["08:35", "A", "V19", "T", "25"],
    ["09:00", "A", "V20", "F", "30"], ["09:10", "A", "V7", "O", "0"],
    ["09:20", "A", "V11", "O", "0"], ["09:30", "A", "V18", "O", "0"],
    ["09:35", "A", "V20", "O", "0"], ["09:50", "A", "V21", "F", "30"],
    ["10:00", "A", "V22", "T", "10"], ["10:05", "C", "V19", "F", "25"],
    ["10:10", "C", "V21", "F", "10"], ["10:20", "C", "V22", "F", "10"],
    ["10:30", "B", "T1", "O", "60"], ["10:50", "B", "F1", "O", "120"],
]

BOOKS = {
    "策略A_优先级": (DOCS / "作业验收预期表_策略A_优先级.xlsx", "PRIORITY"),
    "策略B_时间顺序": (DOCS / "作业验收预期表_策略B_时间顺序.xlsx", "TIME_ORDER"),
}

TOKEN = None


def http(method, path, body=None, token=True):
    data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    headers = {"Accept": "application/json"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    if token and TOKEN:
        headers["Authorization"] = "Bearer " + TOKEN
    request = urllib.request.Request(BASE_URL + path, data=data, headers=headers, method=method)
    with urllib.request.urlopen(request, timeout=TIMEOUT) as response:
        payload = json.loads(response.read().decode("utf-8"))
    if not payload.get("success", False):
        raise RuntimeError(payload.get("message") or "request failed")
    return payload.get("data")


def try_http(method, path, body=None, token=True):
    try:
        return http(method, path, body, token), None
    except Exception as exc:
        return None, exc


def ensure_auth():
    global TOKEN
    data, err = try_http("GET", "/api/admin/snapshot", token=False)
    if err is None:
        return
    if isinstance(err, urllib.error.HTTPError) and err.code in (401, 403):
        username = "verify_admin"
        password = "verify_admin_123456"
        try_http("POST", "/api/auth/register", {"username": username, "password": password, "role": "ADMIN"}, token=False)
        auth = http("POST", "/api/auth/login", {"username": username, "password": password}, token=False)
        TOKEN = auth.get("token")
        http("GET", "/api/admin/snapshot")
        return
    raise RuntimeError(f"无法连接 {BASE_URL}：{err}")


def minutes(hhmm):
    h, m = [int(x) for x in hhmm.split(":")]
    return h * 60 + m


def advance_to(current, target):
    delta = minutes(target) - minutes(current)
    if delta < 0:
        delta += 24 * 60
    if delta:
        http("POST", f"/api/admin/clock/advance?minutes={delta}")
    return target


def run_strategy(strategy):
    http("POST", "/api/admin/config", {
        "waitingCapacity": 10,
        "pileSlotCapacity": 3,
        "fastPileCount": 3,
        "slowPileCount": 2,
        "parkingRate": 0.5,
        "parkingGracePeriodMinutes": 5,
    })
    http("POST", "/api/admin/clock/reset?time=06:00")
    snapshots = []
    current = "06:00"
    for event in EVENTS:
        t, action, subject, mode, amount = event
        if action == "B":
            current = advance_to(current, t)
            path = f"/api/admin/piles/{urllib.parse.quote(subject)}/fault?minutes={urllib.parse.quote(amount)}&strategy={strategy}"
            snapshots.append(http("POST", path))
        else:
            snapshots.append(http("POST", "/api/acceptance/event", {
                "time": t, "action": action, "subject": subject, "mode": mode, "amount": amount,
            }, token=False))
            current = t
    return snapshots, settle_for_details()


def settle_for_details(minutes_to_advance=300):
    http("POST", f"/api/admin/clock/advance?minutes={minutes_to_advance}")
    return http("GET", "/api/admin/snapshot")


def dec(value):
    return Decimal(str(value or 0))


def fmt(value, places=None):
    d = dec(value)
    if places is not None:
        d = d.quantize(Decimal("1." + "0" * places), rounding=ROUND_HALF_UP)
    s = format(d.normalize(), "f")
    return "0" if s == "-0" else s


def rate_at(hhmm):
    t = minutes(hhmm)
    if 600 <= t < 900 or 1080 <= t < 1260:
        return Decimal("1.0")
    if 420 <= t < 600 or 900 <= t < 1080 or 1260 <= t < 1380:
        return Decimal("0.7")
    return Decimal("0.4")


def hhmm(value):
    if not value:
        return None
    text = str(value)
    match = re.search(r"(\d{2}:\d{2})", text)
    return match.group(1) if match else None


def current_fee(req, snap):
    start = hhmm(req.get("startedAt"))
    end = hhmm(snap.get("now"))
    kwh = dec(req.get("chargedKwh"))
    if not start or not end or kwh <= 0:
        return Decimal("0")
    total = max(1, minutes(end) - minutes(start))
    cursor = minutes(start)
    charge_fee = Decimal("0")
    while cursor < minutes(start) + total:
        next_cursor = minutes(start) + total
        for boundary in [420, 600, 900, 1080, 1260, 1380]:
            if cursor < boundary <= minutes(start) + total:
                next_cursor = min(next_cursor, boundary)
                break
        part = Decimal(next_cursor - cursor) / Decimal(total)
        hh = str((cursor // 60) % 24).zfill(2)
        mm = str(cursor % 60).zfill(2)
        charge_fee += kwh * part * rate_at(f"{hh}:{mm}")
        cursor = next_cursor
    return charge_fee + kwh * Decimal("0.8")


def request_by_vehicle(snap):
    result = {}
    for req in snap.get("requests", []):
        prev = result.get(req.get("vehicleId"))
        if not prev or str(req.get("createdAt")) >= str(prev.get("createdAt")):
            result[req.get("vehicleId")] = req
    return result


def pile_cell(snap, pile_id, slot):
    pile = next((p for p in snap.get("piles", []) if p.get("id") == pile_id), None)
    if not pile:
        return ""
    if pile.get("status") == "FAULT":
        return "✕故障"
    if slot == 0:
        vid = pile.get("currentVehicle")
        if not vid:
            return ""
        req = request_by_vehicle(snap).get(vid, {})
        return f"({vid},{fmt(req.get('chargedKwh'), 2)},{fmt(current_fee(req, snap), 2)})"
    queued = pile.get("queuedVehicles") or []
    if len(queued) < slot:
        return ""
    return f"({queued[slot - 1]},0,0)"


def waiting_cell(snap):
    reqs = request_by_vehicle(snap)
    normal = []
    for vid in (snap.get("waitingFast") or []) + (snap.get("waitingSlow") or []):
        req = reqs.get(vid, {})
        mode = "F" if req.get("mode") == "FAST" else "T"
        normal.append(f"({vid},{mode},{fmt(req.get('requestedKwh'))})")
    return " ".join(normal)


def redispatch_cell(snap):
    redispatch = []
    for req in snap.get("requests", []):
        if req.get("status") == "REDISPATCHING":
            mode = "F" if req.get("mode") == "FAST" else "T"
            redispatch.append(f"({req.get('vehicleId')},{mode},{fmt(req.get('requestedKwh'))})")
    return " ".join(redispatch)


def actual_table_rows(snapshots):
    rows = [
        ["", "", "(车号,已充电量,当前费用)", "", "", "", "", "(车号,充电类型,充电量)", "故障后车辆，不属于普通等候区"],
        ["时刻", "事件 (类型,id,模式,值)", "快充1 (F1)", "快充2 (F2)", "快充3 (F3)", "慢充1 (T1)", "慢充2 (T2)", "等候区(10辆)", "故障重调度队列"],
    ]
    pile_ids = ["F1", "F2", "F3", "T1", "T2"]
    for event, snap in zip(EVENTS, snapshots):
        rows.append([event[0], f"({event[1]},{event[2]},{event[3]},{event[4]})", *[pile_cell(snap, p, 0) for p in pile_ids], waiting_cell(snap), redispatch_cell(snap)])
        rows.append(["", "", *[pile_cell(snap, p, 1) for p in pile_ids], "", ""])
        rows.append(["", "", *[pile_cell(snap, p, 2) for p in pile_ids], "", ""])
    return rows


def excel_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["预期表"]
    return [[ws.cell(r, c).value for c in range(1, 10)] for r in range(1, 1 + 2 + len(EVENTS) * 3)]


def normalize_cell(value):
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value).strip())
    text = text.replace("[优先]", "[故障重调度]").replace("[故障优先重调度]", "[故障重调度]")
    text = text.replace("✕故障", "✕故障")
    return text


def compare_rows(expected, actual):
    mismatches = []
    for r, (erow, arow) in enumerate(zip(expected, actual), start=1):
        for c, (e, a) in enumerate(zip(erow, arow), start=1):
            ne, na = normalize_cell(e), normalize_cell(a)
            if ne != na:
                mismatches.append({"row": r, "col": c, "expected": ne, "actual": na})
    return mismatches


def detail_rows_from_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["详单(自动结算)"]
    rows = []
    for r in range(1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, 10)]
        if any(v is not None for v in values):
            rows.append(values)
    return rows


def actual_detail_rows(snap):
    by_request = {r.get("id"): r for r in snap.get("requests", [])}
    rows = []
    for d in snap.get("details", []):
        req = by_request.get(d.get("requestId"), {})
        rows.append([
            d.get("vehicleId"), d.get("pileId"), hhmm(d.get("startedAt")), hhmm(d.get("stoppedAt")),
            fmt(d.get("chargedKwh"), 2), fmt(d.get("chargeFee"), 2), fmt(d.get("serviceFee"), 2), fmt(d.get("totalFee"), 2),
            req.get("status", ""),
        ])
    return rows


def normalize_detail(rows):
    clean = []
    for row in rows:
        if row and str(row[0]).strip() in {"车辆", "车辆编号", "vehicleId", "车号"}:
            continue
        clean.append([normalize_cell(x) for x in row])
    return clean


def compare_details(expected, actual):
    exp = normalize_detail(expected)
    act = normalize_detail(actual)
    max_len = max(len(exp), len(act))
    mismatches = []
    for i in range(max_len):
        erow = exp[i] if i < len(exp) else []
        arow = act[i] if i < len(act) else []
        if erow != arow:
            mismatches.append({"row": i + 1, "expected": erow, "actual": arow})
    return mismatches


def text_warnings(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    warnings = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                text = str(value or "")
                if "优先级等候区" in text:
                    warnings.append(f"{ws.title}: {text[:120]}")
                if "故障后进入普通等候区" in text or "故障车辆进入普通等候区" in text:
                    warnings.append(f"{ws.title}: {text[:120]}")
    return warnings[:10]


def html_table(rows, max_rows=None):
    body = []
    for i, row in enumerate(rows[:max_rows] if max_rows else rows):
        tag = "th" if i == 0 else "td"
        body.append("<tr>" + "".join(f"<{tag}>{escape(normalize_cell(c))}</{tag}>" for c in row) + "</tr>")
    return "<table>" + "".join(body) + "</table>"


def mismatch_table(mismatches, limit=40):
    if not mismatches:
        return "<p class='ok'>无差异</p>"
    rows = [["行", "列", "Excel 预期", "Web 实际"]]
    for m in mismatches[:limit]:
        rows.append([m.get("row"), m.get("col", ""), m.get("expected"), m.get("actual")])
    extra = "" if len(mismatches) <= limit else f"<p>仅显示前 {limit} 条，共 {len(mismatches)} 条。</p>"
    return extra + html_table(rows)


def render_report(results):
    cards = []
    for name, result in results.items():
        ok_table = not result["table_mismatches"]
        ok_detail = not result["detail_mismatches"]
        status = "完全符合" if ok_table and ok_detail and not result["warnings"] else "存在差异/文案风险"
        cards.append(f"""
        <section class='card'>
          <h2>{escape(name)}：{status}</h2>
          <p><b>预期表差异：</b>{len(result['table_mismatches'])}；<b>详单差异：</b>{len(result['detail_mismatches'])}；<b>文案风险：</b>{len(result['warnings'])}</p>
          <h3>预期表差异</h3>{mismatch_table(result['table_mismatches'])}
          <h3>详单差异</h3>{mismatch_table(result['detail_mismatches'])}
          <h3>文案风险</h3>{'<ul>' + ''.join('<li>'+escape(w)+'</li>' for w in result['warnings']) + '</ul>' if result['warnings'] else '<p class="ok">无</p>'}
          <h3>Web 实际最终详单</h3>{html_table([["车辆", "桩", "开始", "结束", "电量", "充电费", "服务费", "总费用", "状态"], *result['actual_details']])}
          <h3>Web 实际最后 6 行预期表镜像</h3>{html_table(result['actual_rows'][-6:])}
        </section>""")
    html = f"""<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><title>Web 运行结果 vs Excel 验收表</title>
    <style>body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:24px;background:#f7f3ec;color:#263238}}.card{{background:#fff;border:1px solid #e7dfd2;border-radius:16px;padding:18px;margin:16px auto;max-width:1400px;box-shadow:0 12px 34px rgba(30,25,20,.08)}}table{{border-collapse:collapse;width:100%;font-size:12px;margin:8px 0 16px}}td,th{{border:1px solid #ded6ca;padding:6px;vertical-align:top}}th{{background:#eef3f5}}.ok{{color:#2b8a3e;font-weight:700}}h1{{max-width:1400px;margin:0 auto 16px}}h2{{margin-top:0}}</style></head><body>
    <h1>Web/API 实际运行结果与两份作业验收预期表对照</h1><p style='max-width:1400px;margin:0 auto'>后端地址：{escape(BASE_URL)}；生成时间：{escape(time.strftime('%Y-%m-%d %H:%M:%S'))}</p>{''.join(cards)}</body></html>"""
    OUT_HTML.write_text(html, encoding="utf-8")


def main():
    ensure_auth()
    results = {}
    for name, (path, strategy) in BOOKS.items():
        if not path.exists():
            raise FileNotFoundError(path)
        snapshots, final_snapshot = run_strategy(strategy)
        actual_rows = actual_table_rows(snapshots)
        expected_rows = excel_rows(path)
        actual_details = actual_detail_rows(final_snapshot)
        expected_details = detail_rows_from_excel(path)
        results[name] = {
            "file": str(path),
            "strategy": strategy,
            "table_mismatches": compare_rows(expected_rows, actual_rows),
            "detail_mismatches": compare_details(expected_details, actual_details),
            "warnings": text_warnings(path),
            "actual_rows": actual_rows,
            "actual_details": actual_details,
        }
    OUT_JSON.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    render_report(results)
    print(json.dumps({
        "html": str(OUT_HTML),
        "json": str(OUT_JSON),
        "summary": {k: {"table_mismatches": len(v["table_mismatches"]), "detail_mismatches": len(v["detail_mismatches"]), "warnings": len(v["warnings"])} for k, v in results.items()},
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
