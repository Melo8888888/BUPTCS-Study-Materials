import openpyxl

def dump_file(filename):
    print(f"=== {filename} ===")
    wb = openpyxl.load_workbook("docs/" + filename, data_only=True)
    ws = wb["预期表"]
    for r in range(80, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, 10)]
        # format values
        row_str = []
        for v in row_vals:
            if v is None:
                row_str.append("")
            else:
                row_str.append(str(v).replace("\u2715", "X"))
        print(f"Row {r:3d}: {row_str}")

dump_file("作业验收预期表_策略A_优先级.xlsx")
dump_file("作业验收预期表_策略B_时间顺序.xlsx")
